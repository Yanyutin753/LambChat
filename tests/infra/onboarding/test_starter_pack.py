import importlib.util
from pathlib import Path


def test_starter_pack_is_available():
    assert importlib.util.find_spec("src.infra.onboarding") is not None, (
        "starter onboarding is missing"
    )


def test_starter_artifact_scripts_generate_files(tmp_path):
    import json
    import subprocess
    import sys
    import zipfile

    root = Path("src/infra/onboarding/skills")
    for skill, suffix, payload in [
        ("documents", "docx", {"title": "研究报告", "paragraphs": ["来源与结论"]}),
        ("spreadsheets", "xlsx", {"headers": ["项目", "金额"], "rows": [["预算", 42]]}),
        (
            "presentations",
            "pptx",
            {"title": "季度回顾", "slides": [{"title": "结果", "bullets": ["交付完成"]}]},
        ),
    ]:
        script = root / f"lambchat-{skill}" / "scripts" / f"create_{skill}.py"
        assert script.exists(), f"{skill} runnable helper is missing"
        source = tmp_path / f"{skill}.json"
        source.write_text(json.dumps(payload), encoding="utf-8")
        target = tmp_path / f"output.{suffix}"
        subprocess.run([sys.executable, str(script), str(source), str(target)], check=True)
        with zipfile.ZipFile(target) as archive:
            assert archive.testzip() is None
            assert "[Content_Types].xml" in archive.namelist()
        if suffix == "xlsx":
            from openpyxl import load_workbook

            book = load_workbook(target)
            assert list(book.active.values) == [("项目", "金额"), ("预算", 42)]
        elif suffix == "pptx":
            from pptx import Presentation

            deck = Presentation(target)
            assert deck.slides[0].shapes.title.text == "结果"
        else:
            from xml.etree import ElementTree as ET

            with zipfile.ZipFile(target) as archive:
                texts = ET.fromstring(archive.read("word/document.xml")).iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                )
                assert [node.text for node in texts] == ["研究报告", "来源与结论"]


class Collection:
    def __init__(self, docs=None):
        self.docs = docs or []
        self.fail_path = None

    async def find_one(self, query, projection=None):
        return next(
            (d.copy() for d in self.docs if all(d.get(k) == v for k, v in query.items())), None
        )

    async def update_one(self, query, update, upsert=False):
        if query.get("file_path") == self.fail_path and self.fail_path:
            raise RuntimeError("Interrupted installation")
        doc = next((d for d in self.docs if all(d.get(k) == v for k, v in query.items())), None)
        if doc is None and upsert:
            doc = {**query, **update.get("$setOnInsert", {})}
            self.docs.append(doc)
        if doc is not None:
            for key, value in update.get("$addToSet", {}).items():
                values = doc.setdefault("metadata", {}).setdefault(key.split(".")[-1], [])
                if value not in values:
                    values.append(value)


import pytest


@pytest.fixture
def stores(monkeypatch):
    from types import SimpleNamespace

    from bson import ObjectId

    from src.infra.onboarding import starter_pack as pack

    uid = ObjectId()
    users = Collection([{"_id": uid, "metadata": {"disabled_skills": ["lambchat-writing"]}}])
    files = Collection()
    roles = Collection()
    monkeypatch.setattr(pack.settings, "ENABLE_SKILLS", True)
    monkeypatch.setattr(pack, "UserStorage", lambda: SimpleNamespace(collection=users))
    monkeypatch.setattr(pack, "PersonaPresetStorage", lambda: SimpleNamespace(collection=roles))

    class Skills:
        def _get_files_collection(self):
            return files

        async def invalidate_user_cache(self, user_id):
            pass

    monkeypatch.setattr(pack, "SkillStorage", Skills)
    return pack, str(uid), users, files, roles


async def test_initialization_preserves_existing_skill_and_disabled_state(stores):
    pack, uid, users, files, _ = stores
    files.docs.append(
        {
            "user_id": uid,
            "skill_name": "lambchat-writing",
            "file_path": "SKILL.md",
            "content": "My own instructions",
        }
    )
    await pack.ensure_user_starter_pack(uid)
    assert len({d["skill_name"] for d in files.docs}) == 6
    writing = [d for d in files.docs if d["skill_name"] == "lambchat-writing"]
    assert len(writing) == 1
    assert writing[0]["content"] == "My own instructions"
    assert users.docs[0]["metadata"]["disabled_skills"] == ["lambchat-writing"]


async def test_completed_installation_does_not_restore_deleted_skill(stores):
    pack, uid, _, files, _ = stores
    await pack.ensure_user_starter_pack(uid)
    files.docs = [d for d in files.docs if d["skill_name"] != "lambchat-documents"]
    await pack.ensure_user_starter_pack(uid)
    assert not any(d["skill_name"] == "lambchat-documents" for d in files.docs)


async def test_interrupted_installation_resumes_without_overwriting_edits(stores):
    pack, uid, _, files, _ = stores
    files.fail_path = "scripts/create_documents.py"
    with pytest.raises(RuntimeError, match="Interrupted"):
        await pack.ensure_user_starter_pack(uid)
    for doc in files.docs:
        if doc["skill_name"] == "lambchat-documents" and doc["file_path"] == "SKILL.md":
            doc["content"] = "Edited during retry"
    files.fail_path = None
    await pack.ensure_user_starter_pack(uid)
    document_files = {
        d["file_path"]: d["content"] for d in files.docs if d["skill_name"] == "lambchat-documents"
    }
    assert document_files["SKILL.md"] == "Edited during retry"
    assert "scripts/create_documents.py" in document_files


async def test_personas_are_public_bound_and_preserve_admin_edits(stores):
    pack, _, _, _, roles = stores
    await pack.ensure_starter_personas()
    roles.docs[0]["system_prompt"] = "Admin edited this role"
    await pack.ensure_starter_personas()
    assert len(roles.docs) == 6
    assert roles.docs[0]["system_prompt"] == "Admin edited this role"
    for role in roles.docs:
        assert role["scope"] == "global" and role["visibility"] == "public"
        assert role["status"] == "published"
        assert role["skill_names"] and all(
            name.startswith("lambchat-") for name in role["skill_names"]
        )
        assert set(role["starter_prompts"][0]["text"]) == {"en", "zh", "ja", "ko", "ru"}


def test_localized_persona_keeps_admin_edits():
    from src.infra.onboarding import starter_pack as pack

    item = pack.load_catalog()[0]
    assert hasattr(pack, "localize_starter_persona"), "starter persona localization missing"
    import hashlib

    from src.kernel.schemas.persona_preset import PersonaPreset

    preset = PersonaPreset(
        id=hashlib.sha256(item["skill_name"].encode()).hexdigest()[:24],
        scope="global",
        name="Research assistant",
        description=item["description"]["en"],
        system_prompt="Research",
        visibility="public",
        status="published",
    )
    localized = pack.localize_starter_persona(preset, "zh-CN")
    assert localized.name == "研究助手"
    assert preset.name == "Research assistant"
    edited = preset.model_copy(update={"name": "My research role"})
    assert pack.localize_starter_persona(edited, "zh").name == "My research role"


async def test_malformed_existing_metadata_is_not_overwritten(stores):
    pack, uid, _, files, _ = stores
    files.docs.append(
        {
            "user_id": uid,
            "skill_name": "lambchat-writing",
            "file_path": "__meta__",
            "content": "not JSON",
        }
    )
    await pack.ensure_user_starter_pack(uid)
    assert len({d["skill_name"] for d in files.docs}) == 6
    assert (
        next(d for d in files.docs if d["skill_name"] == "lambchat-writing")["content"]
        == "not JSON"
    )


def test_all_user_creation_paths_share_starter_initialization():
    import inspect

    from src.infra.skill.manager import SkillManager
    from src.infra.user.manager import UserManager
    from src.infra.user.storage import UserStorage

    assert "initialize_user_starter_pack" in inspect.getsource(UserStorage.create)
    assert "initialize_user_starter_pack" in inspect.getsource(UserManager.get_user)
    assert "initialize_user_starter_pack" in inspect.getsource(SkillManager.get_effective_skills)


def test_starter_skill_description_uses_requested_language():
    from src.infra.onboarding import starter_pack as pack
    from src.infra.skill.types import UserSkill

    assert hasattr(pack, "localize_starter_skill"), "skill localization is missing"
    skill = UserSkill(
        skill_name="lambchat-research",
        description="Find sources, verify claims and write a cited research report.",
    )
    assert (
        pack.localize_starter_skill(skill, "ja").description
        == "情報を検索し、事実を確認して出典付きの調査報告を作成します。"
    )
    custom = skill.model_copy(update={"description": "My custom instructions"})
    assert pack.localize_starter_skill(custom, "zh").description == "My custom instructions"


def test_localized_snapshot_uses_the_same_role_name():
    import hashlib

    from src.infra.onboarding import starter_pack as pack
    from src.kernel.schemas.persona_preset import PersonaPresetSnapshot

    snapshot = PersonaPresetSnapshot(
        preset_id=hashlib.sha256(b"lambchat-research").hexdigest()[:24],
        name="Research assistant",
        system_prompt="Research",
        skill_names=["lambchat-research"],
        missing_skill_names=[],
        version=1,
    )
    assert pack.localize_starter_persona(snapshot, "ko").name == "연구 도우미"


async def test_research_and_writing_roles_can_also_deliver_document_files(stores):
    pack, _, _, _, roles = stores
    await pack.ensure_starter_personas()
    research = next(role for role in roles.docs if role["name"] == "Research assistant")
    writer = next(role for role in roles.docs if role["name"] == "Writing assistant")
    assert set(research["skill_names"]) == {
        "lambchat-research",
        "lambchat-writing",
        "lambchat-documents",
    }
    assert set(writer["skill_names"]) == {"lambchat-writing", "lambchat-documents"}


async def test_missing_skill_bundle_remains_retryable(stores, monkeypatch, tmp_path):
    pack, uid, users, _, _ = stores
    monkeypatch.setattr(pack, "ROOT", tmp_path)
    with pytest.raises(ValueError, match="SKILL.md"):
        await pack.ensure_user_starter_pack(uid)
    assert not users.docs[0]["metadata"].get("starter_skills")


def test_spreadsheet_preserves_formula_like_input_as_text(tmp_path):
    import runpy

    from openpyxl import load_workbook

    script = Path(
        "src/infra/onboarding/skills/lambchat-spreadsheets/scripts/create_spreadsheets.py"
    )
    create = runpy.run_path(str(script))["create"]
    output = tmp_path / "data.xlsx"
    create({"headers": ["value"], "rows": [["=1+1"], ["@SUM(A1)"], ["-123"]]}, output)
    sheet = load_workbook(output).active
    assert [sheet.cell(row, 1).data_type for row in range(2, 5)] == ["s", "s", "s"]
    assert sheet["A2"].value == "=1+1"


def test_invalid_spreadsheet_rows_do_not_replace_existing_file(tmp_path):
    import runpy

    script = Path(
        "src/infra/onboarding/skills/lambchat-spreadsheets/scripts/create_spreadsheets.py"
    )
    output = tmp_path / "data.xlsx"
    output.write_bytes(b"existing user file")
    with pytest.raises(ValueError, match="header width"):
        runpy.run_path(str(script))["create"](
            {"headers": ["one"], "rows": [["one", "two"]]}, output
        )
    assert output.read_bytes() == b"existing user file"


def test_starter_personas_use_existing_avatar_icon_format():
    from src.infra.onboarding.starter_pack import load_catalog

    allowed = {"icon:academic", "icon:writing", "icon:data", "icon:productivity", "icon:coding"}
    assert all(item["icon"] in allowed for item in load_catalog())


def test_direct_skill_listing_and_persona_activation_initialize_existing_users():
    import inspect

    from src.api.routes.persona_preset import use_persona_preset
    from src.api.routes.skill import list_user_skills

    assert "initialize_user_starter_pack" in inspect.getsource(list_user_skills)
    assert "initialize_user_starter_pack" in inspect.getsource(use_persona_preset)


def test_persona_search_includes_translated_starter_names():
    from src.infra.persona_preset.storage import PersonaPresetStorage

    query = PersonaPresetStorage()._build_visible_query(user_id="user-1", q="研究助手")
    assert "starter_search_text" in str(query)


@pytest.mark.parametrize("text", ["연구 도우미", "Помощник исследователя"])
def test_localized_persona_search_accepts_korean_and_russian(text):
    from src.infra.persona_preset.storage import _build_persona_search_terms

    assert _build_persona_search_terms(text), (
        "valid localized names must not produce an empty query"
    )
