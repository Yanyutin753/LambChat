import io
import zipfile

import pytest
from PIL import Image

from src.api.routes.upload_content_cover import preview_text, render_content_cover


def test_text_preview_is_real_bounded_content():
    assert preview_text(b"actual_col,actual_value\nhello,7\n", "csv").startswith(
        "actual_col,actual_value"
    )
    assert len(preview_text(b"x" * 100000, "txt")) <= 6000
    image = Image.open(io.BytesIO(render_content_cover(b"hello world", "txt")))
    assert image.size == (560, 315)


def test_archive_preview_lists_without_extracting_paths():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("../unsafe.txt", "not executed")
        archive.writestr("report.csv", "a,b")
    text = preview_text(data.getvalue(), "zip")
    assert "report.csv" in text and "not executed" not in text


def test_office_zip_bomb_is_rejected_before_conversion():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (8 * 1024 * 1024))
    with pytest.raises(ValueError):
        preview_text(data.getvalue(), "docx")


def test_office_preview_reads_actual_presentation_text():
    from pptx import Presentation
    from pptx.util import Inches

    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text = "Actual quarterly result"
    data = io.BytesIO()
    document.save(data)
    assert "Actual quarterly result" in preview_text(data.getvalue(), "pptx")


def test_word_preview_reads_actual_document_text():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        )
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Actual word report</w:t></w:r></w:p></w:body></w:document>',
        )
    assert "Actual word report" in preview_text(data.getvalue(), "docx")
