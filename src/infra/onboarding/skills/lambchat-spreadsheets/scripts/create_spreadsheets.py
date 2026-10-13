"""Create an XLSX containing literal tabular data and an optional bar chart."""

import json
import sys
from pathlib import Path


def create(data, output):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font

    headers, rows = data["headers"], data["rows"]
    if not headers or not isinstance(headers, list) or not isinstance(rows, list):
        raise ValueError("Expected non-empty headers and a rows list")
    if any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
        raise ValueError("Every row must match the header width")
    book = Workbook()
    sheet = book.active
    sheet.title = data.get("sheet_name", "Data")
    for values in [headers, *rows]:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if data.get("chart"):
        spec = data["chart"]
        category, value = spec["category_column"], spec["value_column"]
        if not rows or not (1 <= category <= len(headers) and 1 <= value <= len(headers)):
            raise ValueError("Chart columns must exist and data rows must be non-empty")
        chart = BarChart()
        chart.title = spec.get("title", "Summary")
        chart.add_data(
            Reference(sheet, min_col=value, max_col=value, min_row=1, max_row=sheet.max_row),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(sheet, min_col=category, max_col=category, min_row=2, max_row=sheet.max_row)
        )
        sheet.add_chart(chart, "H2")
    book.save(output)


if __name__ == "__main__":
    create(json.loads(Path(sys.argv[1]).read_text(encoding="utf-8")), Path(sys.argv[2]))
