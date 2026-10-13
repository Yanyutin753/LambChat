"""Create an editable widescreen presentation from UTF-8 JSON."""

import json
import sys
from pathlib import Path


def create(data, output):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    slides = data["slides"]
    if not isinstance(slides, list) or not slides:
        raise ValueError("Expected a non-empty slides list")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = data.get("title", "Presentation")
    for item in slides:
        if (
            not isinstance(item["title"], str)
            or not isinstance(item["bullets"], list)
            or not all(isinstance(v, str) for v in item["bullets"])
        ):
            raise ValueError("Each slide needs a string title and string bullets")
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = item["title"]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, text in enumerate(item["bullets"]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.font.size = Pt(24)
    deck.save(output)


if __name__ == "__main__":
    create(json.loads(Path(sys.argv[1]).read_text(encoding="utf-8")), Path(sys.argv[2]))
